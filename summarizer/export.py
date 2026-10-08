import io

from docx import Document


def summary_to_docx(title, summary, keywords=None, meta=None) -> bytes:
    doc = Document()
    doc.add_heading("Document Summary", 0)
    doc.add_paragraph(f"Source: {title}")
    if meta:
        doc.add_paragraph(meta)
    doc.add_heading("Summary", 1)
    doc.add_paragraph(summary)
    if keywords:
        doc.add_heading("Keywords", 1)
        doc.add_paragraph(", ".join(keywords))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()